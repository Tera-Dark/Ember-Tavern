from pathlib import Path
import re
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

root=Path(__file__).resolve().parent.parent
source=(root/'docs/PRD.md').read_text(encoding='utf-8')
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(1.9);section.bottom_margin=Cm(1.9)
section.left_margin=Cm(2);section.right_margin=Cm(2)
section.header_distance=Cm(.8);section.footer_distance=Cm(.8)
section.different_first_page_header_footer=True
styles=doc.styles
normal=styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string('303C32')
normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.4
for level,size in [(1,18),(2,14),(3,12)]:
 s=styles[f'Heading {level}'];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string('294A37');s.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
 s.paragraph_format.space_before=Pt(20 if level==1 else 15);s.paragraph_format.space_after=Pt(9);s.paragraph_format.keep_with_next=True
if 'Code Block' not in styles:
 s=styles.add_style('Code Block',WD_STYLE_TYPE.PARAGRAPH);s.font.name='Consolas';s.font.size=Pt(8.5);s.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');s.paragraph_format.line_spacing=1.25;s.paragraph_format.space_after=Pt(2);s.paragraph_format.left_indent=Cm(.2)

header=section.header.paragraphs[0]
header.add_run('余烬酒馆  /  EMBER TAVERN').font.color.rgb=RGBColor.from_string('8F7546')
header.runs[0].font.size=Pt(8)
header.add_run('                                       产品需求文档 · 2.1.0-beta.2 测试版').font.size=Pt(8)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('2026.09.30  ·  余烬酒馆  |  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('7A877A')
r=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)

# Cover
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(50)
r=p.add_run('PRODUCT REQUIREMENTS  /  2.1.0-beta.2 测试版');r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('9E7D40')
p=doc.add_paragraph();p.paragraph_format.space_after=Pt(9)
r=p.add_run('余烬酒馆');r.font.name='Calibri';r.font.size=Pt(36);r.font.color.rgb=RGBColor.from_string('253E2D');r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
p=doc.add_paragraph('小宿主 · 可插拔多人跑团');p.runs[0].font.size=Pt(19);p.runs[0].font.color.rgb=RGBColor.from_string('62735F');p.paragraph_format.space_after=Pt(25)
im=Image.open(root/'static/harbor.jpg').convert('RGB');w,h=im.size;target_h=int(w/2.9);top=max(0,(h-target_h)//2);im=im.crop((0,top,w,top+target_h));cover=root/'docs/prd-cover.jpg';im.save(cover,quality=90)
doc.add_picture(str(cover),width=Cm(17))
p=doc.add_paragraph('让一位房主和一桌玩家，走进同一个世界。');p.paragraph_format.space_before=Pt(24);p.paragraph_format.space_after=Pt(13);p.runs[0].font.size=Pt(15);p.runs[0].font.color.rgb=RGBColor.from_string('294A37')
p=doc.add_paragraph('内核守正史，插件做演绎。地图、图标、台本与声音按需组装，关掉也能继续冒险。');p.paragraph_format.space_after=Pt(24)
for text in ['日期  /  2026-09-30','范围  /  模块化全栈 MVP + Plugin SDK v1 + GitHub 发布模板','状态  /  51 项本地测试通过；真实供应商待配置，GitHub 状态以 Actions 为准']:
 p=doc.add_paragraph(text);p.paragraph_format.space_after=Pt(9);p.runs[0].font.size=Pt(10);p.runs[0].font.color.rgb=RGBColor.from_string('7D897A')
doc.add_page_break()

def shade(p,color):
 pr=p._p.get_or_add_pPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);pr.append(sh)

def inline(p,text):
 parts=re.split(r'(\*\*.*?\*\*|`[^`]*`|\[[^\]]*\]\([^)]+\))',text)
 for item in parts:
  if not item:continue
  if item.startswith('**') and item.endswith('**'):p.add_run(item[2:-2]).bold=True
  elif item.startswith('`') and item.endswith('`'):
   r=p.add_run(item[1:-1]);r.font.name='Consolas';r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('8F6A2F')
  elif re.match(r'^\[.*\]\(.*\)$',item):
   m=re.match(r'^\[(.*?)\]\((.*?)\)$',item);link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(m[2],RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');prop=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'8F6A2F');prop.append(c);u=OxmlElement('w:u');u.set(qn('w:val'),'single');prop.append(u);r.append(prop);t=OxmlElement('w:t');t.text=m[1];r.append(t);link.append(r);p._p.append(link)
  else:p.add_run(item)

lines=source.splitlines();i=1
while i<len(lines):
 line=lines[i].strip()
 if not line or line=='---':i+=1;continue
 if line.startswith('!['):
  m=re.match(r'^!\[(.*?)\]\((.*?)\)$',line)
  if m and (root/'docs'/m[2]).exists():doc.add_picture(str(root/'docs'/m[2]),width=Cm(17))
  i+=1;continue
 if line.startswith('```'):
  code=[];i+=1
  while i<len(lines) and not lines[i].strip().startswith('```'):code.append(lines[i]);i+=1
  for s in code:
   p=doc.add_paragraph(s,style='Code Block');shade(p,'F1F4EC')
  doc.add_paragraph().paragraph_format.space_after=Pt(3)
  i+=1;continue
 if line.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch(r'[:\-\s]+',c) for c in row):rows.append(row)
   i+=1
  if not rows:continue
  table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid';table.autofit=False
  for n,row in enumerate(rows):
   cells=table.add_row().cells
   for idx,text in enumerate(row):
    if idx>=len(cells):break
    cell=cells[idx];cell.width=Cm(17/len(cells));p=cell.paragraphs[0];inline(p,text);p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.3
    for r in p.runs:r.font.size=Pt(8.5)
    if n==0:
     pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'294A37');pr.append(sh)
     for r in p.runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    elif n%2:
     pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F3F5EF');pr.append(sh)
   if n==0:
    pr=table.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');pr.append(rep)
  doc.add_paragraph().paragraph_format.space_after=Pt(2)
  continue
 if line.startswith('#'):
  m=re.match(r'^(#+)\s+(.*)',line);p=doc.add_paragraph(style=f'Heading {min(len(m[1])-1,3)}' if len(m[1])>1 else 'Heading 1');inline(p,m[2]);i+=1;continue
 if line.startswith('>'):
  p=doc.add_paragraph();inline(p,line[1:].strip());shade(p,'EEF2E8');p.paragraph_format.left_indent=Cm(.3);p.paragraph_format.right_indent=Cm(.2)
 elif line.startswith('- '):
  p=doc.add_paragraph(style='List Bullet');inline(p,line[2:])
 elif re.match(r'^\d+\. ',line):
  p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(.25);inline(p,line)
 else:
  p=doc.add_paragraph();inline(p,line)
 i+=1

doc.core_properties.title='余烬酒馆 · 产品需求文档 PRD 2.1.0-beta.2'
doc.core_properties.subject='模块化多人跑团：权限、资源契约、地图、台本、TTS、隔离与回档'
doc.core_properties.author='余烬酒馆项目'
doc.core_properties.keywords='PRD,TRPG,AI,多人跑团,MVP'
output=root/'docs/PRD.docx';doc.save(output)
# Validate the deliverable as a real OOXML document.
checked=Document(output)
print(f'DOCX validated: {len(checked.paragraphs)} paragraphs, {len(checked.tables)} tables, {output.stat().st_size} bytes')
